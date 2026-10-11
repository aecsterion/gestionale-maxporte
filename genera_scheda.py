#!/usr/bin/env python3
"""
Scheda di produzione + packing list (una per ordine).

Uso: python3 genera_scheda.py payload.json out.pdf

Riempie template_produzione.xlsx (fogli SCHEDA_PRODUZIONE e PACKING_LIST):
- testata (righe 1-6) uguale sui due fogli, ripetuta su ogni pagina (righe titolo)
- righe posizione dinamiche dalla riga 9 (stile copiato dalla riga 9 del template)
- packing list a tre blocchi dalla riga 11 (stile copiato dalla riga 11)
- barcode Code128 del numero ordine + logo come immagini
poi converte in PDF con LibreOffice.
"""
import json, sys, os, io, copy, math, shutil, subprocess, tempfile
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.cell_range import CellRange

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, 'template_produzione.xlsx')
LOGO = os.path.join(BASE, 'logo-maxporte.png')

# ── colonne riga posizione (cella iniziale del merge → campo payload) ──
POS_COLS = [
    ('A', 'pos'), ('C', 'modello'), ('K', 'finitura'), ('S', 'l'), ('U', 'h'), ('W', 's'),
    ('Y', 'senso'), ('AA', 'qta'), ('AC', 'tipo'), ('AG', 'serratura'), ('AJ', 'ferramenta'),
    ('AM', 'inserto'), ('AP', 'telaio'), ('AS', 'finitura_telaio'), ('BA', 'coprifili'), ('BD', 'note'),
]
POS_ROW = 9            # riga dati modello nel template
POS_ROW_H = 19.75
LAST_COL = 61          # BI

# packing list: blocchi (col. inizio, col. fine, mappa cella→campo)
PK_ROW = 11
PK_BLOCKS = {
    'minuteria':  (1, 20,  [('A', 'articolo'), ('L', 'um'), ('O', 'qta')]),
    'commessa':   (21, 41, [('U', 'articolo'), ('AG', 'um'), ('AJ', 'qta')]),
    'ubicazioni': (42, 61, [('AP', 'articolo'), ('BC', 'ubicazione')]),
}

HEADER_MAP = {
    '*DATA_STAMPA*': 'data_stampa', '*RAGIONE SOCIALE*': 'cliente',
    '*RIFERIMENTO_CLIENTE*': 'riferimento', '*SETTIMANA_LAVORAZIONE*': 'sett_lavorazione',
    '*SETTIMANA_APPRONTAMENTO*': 'sett_approntamento', '*CODICE_ORDINE_ABBREVIATO*': 'ordine_abbr',
    '*TRASPORTO*': 'trasporto', '*POSA_SI_O_NO*': 'posa', '*ZONA_CONSEGNA*': 'zona', '*NOTE_ORDINE*': 'note',
}


def s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def merges_in_row(ws, row):
    return [CellRange(str(m)) for m in ws.merged_cells.ranges if m.min_row == row and m.max_row == row]


def ensure_row_like(ws, src_row, dst_row, col_from=1, col_to=LAST_COL, height=None):
    """Copia stile celle + merge della riga src su dst (limitato alle colonne indicate)."""
    if src_row == dst_row:
        return
    for c in range(col_from, col_to + 1):
        sc = ws.cell(src_row, c); dc = ws.cell(dst_row, c)
        if sc.has_style:
            dc._style = copy.copy(sc._style)
    existing = {(m.min_col, m.max_col) for m in merges_in_row(ws, dst_row)}
    for m in merges_in_row(ws, src_row):
        if m.min_col >= col_from and m.max_col <= col_to and (m.min_col, m.max_col) not in existing:
            ws.merge_cells(start_row=dst_row, start_column=m.min_col, end_row=dst_row, end_column=m.max_col)
    if height:
        ws.row_dimensions[dst_row].height = height
    elif ws.row_dimensions[src_row].height:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def fill_header(ws, t):
    for row in ws.iter_rows(min_row=1, max_row=7):
        for c in row:
            if isinstance(c.value, str) and c.value in HEADER_MAP:
                key = HEADER_MAP[c.value]
                c.value = s(t.get(key))
                if key not in ('ordine_abbr',):
                    al = copy.copy(c.alignment); al.shrink_to_fit = True; al.wrap_text = False
                    c.alignment = al
    # il placeholder barcode viene sostituito dall'immagine
    for row in ws.iter_rows(min_row=1, max_row=7):
        for c in row:
            if c.value == '*BARCODE_ORDINE*':
                c.value = None


def fix_footer(ws):
    for hf in (ws.oddFooter, ws.evenFooter, ws.firstFooter):
        for part in (hf.left, hf.center, hf.right):
            if part.text and '*' in part.text:
                part.text = part.text.replace('*NUMERO_PAGINA*', '&P').replace('*TOTALE_PAGINE*', '&N')


def make_barcode_png(code):
    try:
        import barcode
        from barcode.writer import ImageWriter
    except Exception:
        return None
    buf = io.BytesIO()
    bc = barcode.get('code128', code, writer=ImageWriter())
    bc.write(buf, options={'module_height': 9.0, 'module_width': 0.28, 'quiet_zone': 2.0,
                           'font_size': 7, 'text_distance': 3.5, 'dpi': 300, 'write_text': True})
    buf.seek(0)
    return buf


def add_images(ws, numero):
    # logo nell'area A1:E6
    if os.path.exists(LOGO):
        img = XLImage(LOGO)
        h = 82
        img.height = h; img.width = int(h * 5244 / 5559)
        ws.add_image(img, 'B1')
    # barcode nell'area AI4:AR6
    png = make_barcode_png(numero) if numero else None
    if png:
        from PIL import Image as PILImage
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        im = PILImage.open(png); w, hh = im.size; png.seek(0)
        img = XLImage(png)
        # area AI4:AR6 = righe 4-6 (~45px): il barcode resta DENTRO i bordi con margine sopra/sotto
        # area AI4:AR6: 10 colonne x (13.5+7+13.5 pt) -> ~234 x 45 px; barcode centrato con margine
        area_w, area_h = 10 * COL_PX, 34 * 96 / 72
        h_px = 36
        w_px = min(int(h_px * w / hh * 1.35), int(area_w - 16))
        img.height = h_px; img.width = w_px
        col_off_tot = max(0, (area_w - w_px) / 2 - 8)   # -8: correzione misurata sul PDF reale
        col = column_index_from_string('AI') - 1 + int(col_off_tot // COL_PX)
        col_off = int(col_off_tot % COL_PX)
        row_off = int((area_h - h_px) / 2)
        marker = AnchorMarker(col=col, colOff=pixels_to_EMU(col_off), row=3, rowOff=pixels_to_EMU(row_off))
        img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(h_px)))
        ws.add_image(img)


COL_PX = 23.4   # larghezza resa di una colonna del template (2.36 caratteri)

def _chars_per_line(ws, row, col, cell):
    """Caratteri per riga stimati per la cella (larghezza del merge / larghezza media carattere)."""
    ci = column_index_from_string(col)
    span = 1
    for m in ws.merged_cells.ranges:
        if m.min_row == row and m.min_col == ci:
            span = m.max_col - m.min_col + 1
            break
    sz = (cell.font.sz or 10)
    char_px = sz * (0.62 if cell.font.b else 0.58)   # Calibri/Carlito, media prudente
    return max(4, int(span * COL_PX / char_px))


def est_lines(text, chars):
    if not text:
        return 1
    return max(1, sum(max(1, math.ceil(len(part) / chars)) for part in str(text).split('\n')))


def build_scheda(ws, t, posizioni):
    fill_header(ws, t)
    if isinstance(ws['AM8'].value, str) and ws['AM8'].value.startswith('*'):
        ws['AM8'].value = 'Inserto / Bugna / Pietra'
    fix_footer(ws)
    add_images(ws, t.get('ordine_numero'))
    n = max(1, len(posizioni))
    for i in range(n):
        r = POS_ROW + i
        ensure_row_like(ws, POS_ROW, r)
        p = posizioni[i] if i < len(posizioni) else {}
        # Allineamenti e "testo a capo" li decide il TEMPLATE: qui si scrivono solo i valori
        # e si calcola l'altezza riga in base alle celle che nel template vanno a capo.
        lines = 1
        for col, key in POS_COLS:
            c = ws[f'{col}{r}']
            c.value = s(p.get(key))
            if c.alignment.wrap_text:
                lines = max(lines, est_lines(c.value, _chars_per_line(ws, POS_ROW, col, c)))
        ws.row_dimensions[r].height = max(POS_ROW_H, 12.5 * lines + 5)
    last = POS_ROW + n - 1
    ws.print_area = f'A1:BI{last}'
    ws.print_title_rows = '1:8'


def build_packing(ws, t, pk):
    fill_header(ws, t)
    fix_footer(ws)
    add_images(ws, t.get('ordine_numero'))
    lens = {k: len(pk.get(k) or []) for k in PK_BLOCKS}
    n = max(1, max(lens.values()))
    for i in range(n):
        r = PK_ROW + i
        row_lines = 1
        for key, (c1, c2, cols) in PK_BLOCKS.items():
            items = pk.get(key) or []
            if i < len(items):
                ensure_row_like(ws, PK_ROW, r, c1, c2)
                for col, field in cols:
                    c = ws[f'{col}{r}']
                    c.value = s(items[i].get(field))
                    # allineamento/a capo decisi dal template; altezza riga in base alle celle a capo
                    if c.alignment.wrap_text:
                        row_lines = max(row_lines, est_lines(c.value, _chars_per_line(ws, PK_ROW, col, c)))
            else:
                # blocco senza voci su questa riga: niente testo né casella, resta solo la riga divisoria
                for col, field in cols:
                    ws[f'{col}{r}'].value = None
                if r == PK_ROW:
                    for c in range(c1, c2 + 1):
                        cell = ws.cell(r, c)
                        b = cell.border
                        if b.left.style == 'medium' and b.right.style == 'medium':
                            cell.border = openpyxl.styles.Border()
                    if key != 'minuteria':
                        pass
        base_h = ws.row_dimensions[PK_ROW].height or 13.5
        ws.row_dimensions[r].height = max(base_h, 12 * row_lines + 4)
        # mantiene la linea verticale di separazione tra i blocchi
        for col in ('T', 'AO'):
            src = ws[f'{col}{PK_ROW}']
            if src.has_style:
                ws[f'{col}{r}']._style = copy.copy(src._style)
        for col in ('U', 'AP'):
            src = ws[f'{col}{PK_ROW}']
            ws[f'{col}{r}'].border = copy.copy(src.border)
    last = max(PK_ROW + n - 1, 40)
    ws.print_area = f'A1:BI{last}'
    ws.print_title_rows = '1:10'


def main():
    src, out = sys.argv[1], sys.argv[2]
    with open(src) as f:
        data = json.load(f)
    t = data.get('testata', {})
    wb = openpyxl.load_workbook(TEMPLATE)
    build_scheda(wb['SCHEDA_PRODUZIONE'], t, data.get('posizioni') or [])
    build_packing(wb['PACKING_LIST'], t, data.get('packing') or {})

    tmpdir = tempfile.mkdtemp()
    try:
        xlsx = os.path.join(tmpdir, 'scheda.xlsx')
        wb.save(xlsx)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmpdir, xlsx],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        shutil.copy(os.path.join(tmpdir, 'scheda.pdf'), out)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == '__main__':
    main()
